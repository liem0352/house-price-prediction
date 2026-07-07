# -*- coding: utf-8 -*-
"""读取考试报告 docx 文件的内容并打印"""
import sys

# 引入 python-docx，用于解析 Word 文档
try:
    import docx
except ImportError:
    print("未安装 python-docx，尝试安装...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

# 打开目标文档
doc_path = r"d:\文件\工作 作业\深度学习\期末\考试报告-.docx"
doc = docx.Document(doc_path)

print("=" * 60)
print("【段落内容】")
print("=" * 60)
for i, p in enumerate(doc.paragraphs):
    style = p.style.name if p.style else "Normal"
    print(f"[{i:03d}] ({style}) {p.text}")

print()
print("=" * 60)
print("【表格内容】")
print("=" * 60)
for ti, table in enumerate(doc.tables):
    print(f"--- 表格 {ti} ---")
    for ri, row in enumerate(table.rows):
        cells_text = " | ".join(cell.text.strip().replace("\n", " / ") for cell in row.cells)
        print(f"  行{ri}: {cells_text}")

print()
print("=" * 60)
print("【文档核心属性】")
print("=" * 60)
cp = doc.core_properties
print(f"标题: {cp.title}")
print(f"作者: {cp.author}")
print(f"创建时间: {cp.created}")
print(f"修改时间: {cp.modified}")
