# -*- coding: utf-8 -*-
"""
按 cell index 直接填充，避免 find_row 漏匹配。
体会部分：拟人化、少字。
"""
from docx import Document
from docx.shared import Pt

DOC = r"d:\文件\工作 作业\深度学习\期末\考试报告-.docx"
doc = Document(DOC)
tbl = doc.tables[0]


def clear_cell(cell):
    """彻底清空单元格：删掉所有段落与表格"""
    tc = cell._tc
    # 删掉所有 w:p
    for p in tc.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
        tc.remove(p)
    for t in tc.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl"):
        tc.remove(t)


def fill(cell, lines, font_size=10.5, default_bold=False):
    """在空 cell 中按 lines 列表写文字"""
    clear_cell(cell)
    for i, item in enumerate(lines):
        if isinstance(item, tuple):
            text, bold = item
        else:
            text, bold = item, default_bold
        p = cell.add_paragraph()
        run = p.add_run(text)
        run.font.size = Pt(font_size)
        run.bold = bold


# ========== Row 0 项目名称 ==========
fill(tbl.cell(0, 1), [("神经网络模型实现 —— 房价预测", True)])

# ========== Row 1 时间 ==========
fill(tbl.cell(1, 1), [("2026 年 6 月 18 日（星期 四） 第 5 节", False)])

# ========== Row 2 分组情况 ==========
fill(tbl.cell(2, 1), [("1 人/组", False)])
fill(tbl.cell(2, 3), [("（独立完成）", False)])

# ========== Row 3 程序运行内容 ==========
fill(tbl.cell(3, 1), [
    ("提交项目源代码文件：", True),
    ("1) 神经网络模型实现_房价预测_完整版.ipynb", False),
    ("2) 任务一_环境搭建与数据集理解.ipynb", False),
    ("3) 任务二_神经网络模型实现_房价预测.ipynb", False),
    ("4) house_price_model.keras", False),
    ("5) metrics.json", False),
    ("6) figures/（4 张训练图表）", False),
])

# ========== Row 4 实现需求目的与要求 ==========
fill(tbl.cell(4, 1), [
    ("考核目的：", True),
    ("掌握神经网络前向/反向传播、数据制备等工作原理；", False),
    ("掌握 TensorFlow 框架的编码特点及应用；", False),
    ("理解模型训练与评估方法；", False),
    ("学会看预测结果及参数调优的可视化图表。", False),
    ("", False),
    ("项目需求：", True),
    ("某政府收集了多家房产公司的二手房销售数据（1460 条 × 81 列），",
     False),
    ("需用深度学习（全连接网络）根据 79 项房屋特征自动预测房价，",
     False),
    ("为政府二手房估价提供价格区间参考。", False),
    ("", False),
    ("评分覆盖：", True),
    ("一、数据集准备与需求分析（20 分）", False),
    ("二、全连接神经网络模型搭建（30 分）", False),
    ("三、模型训练与评估（35 分）", False),
    ("四、任务总结与优化思路（15 分）", False),
])

# ========== Row 5 仪器设备 ==========
fill(tbl.cell(5, 1), [
    ("硬件：Windows 11 64 位 PC", False),
    ("软件：", True),
    ("Anaconda、Python 3.13、JupyterLab", False),
    ("TensorFlow 2.x、Keras", False),
    ("scikit-learn、pandas、numpy、matplotlib", False),
])

# ========== Row 6 实现步骤及数据记录 ==========
fill(tbl.cell(6, 1), [
    ("【任务一：数据集准备与需求分析】", True),
    ("1.1 数据理解：read_csv 加载 1460×81；Id 序号、SalePrice 目标；"
     "36 数值 + 43 分类；范围 $34.9k–$755k，中位 $163k。", False),
    ("", False),
    ("1.2 预处理：", True),
    ("  1.2.1 特征：删 Id；数值中位数填充分类 'None'填充；"
     "StandardScaler；get_dummies 独热；得 303 维。", False),
    ("  1.2.2 标签：右偏（偏度 1.88）→ log1p → 偏度 0.12；"
     "预测后 expm1 还原。", False),
    ("", False),
    ("1.3 问题辨析：回归问题；评估 MAE/RMSE/R²；"
     "1460 行小数据，全连接即可；8:2 切分，random_state=42。", False),
    ("", False),
    ("【任务二：神经网络搭建与训练】", True),
    ("2.1 模型结构：", True),
    ("  输入 Input(303) → Dense(64, ReLU, L2, Drop0.2) → "
     "Dense(32, ...) → Dense(16, ...) → Dense(1)。", False),
    ("2.2 损失/优化器：MSE + MAE；Adam(1e-3, clipnorm=1.0)；"
     "ReduceLROnPlateau。", False),
    ("3.1 训练：batch=32, epochs=500, val=0.2, EarlyStop(patience=80)；"
     "实际跑 194 轮停。", False),
    ("3.2 防过拟合：L2 + Dropout + 早停 + 学习率衰减。", False),
    ("3.3 测试（真实价格）：RMSE=$35,837, MAE=$25,543, R²=0.8326。",
     False),
    ("", False),
    ("【任务四：总结与优化】", True),
    ("4.1 总结：跑通完整流程；log 变换+三重正则有效；表格数据深度学习"
     "R²=0.83 与线性模型相当。", False),
    ("4.2 优化：特征交叉/目标编码降维；改用 XGBoost/LightGBM（可至 "
     "0.88~0.92）；KerasTuner 调参；Flask 部署。", False),
])

# ========== Row 7 结果分析及体会（拟人化·少字） ==========
fill(tbl.cell(7, 1), [
    ("结果：", True),
    ("测试 R²=0.8326，MAE=$25,543，预测值与真实值散点基本沿对角线分布，"
     "残差均值≈0、分布近似正态，模型可用。", False),
    ("", False),
    ("体会：", True),
    ("做之前觉得房价预测就是套模型，跑起来才发现坑不少。", False),
    ("一是标签右偏，不 log 一下 MAE 死活降不下来；", False),
    ("二是特征杂，缺值策略要分类型处理，不能一刀切；", False),
    ("三是小数据集上模型不能堆太大，加 L2、Dropout、早停三件套才稳住。",
     False),
    ("这次最大的收获是理解了'预处理 > 调参'，", False),
    ("也体会到表格数据深度学习并不天然优于线性模型，", False),
    ("未来打算换成 XGBoost 试试。", False),
])

# Row 8 教师评分 留空

doc.save(DOC)
print("已完整填写 考试报告-.docx")
